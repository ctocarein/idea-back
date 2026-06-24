"""Sprint 4 bout en bout : comités (lecture) + upload/parsing d'un deck."""

from __future__ import annotations

from io import BytesIO

import pytest

pytestmark = pytest.mark.integration

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


async def _register(client, email: str) -> dict:
    r = await client.post(
        "/api/v1/auth/register",
        json={"name": "Sophie", "email": email, "password": "s3cret-pwd", "consent": True},
    )
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _make_pptx(titles: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for t in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = t
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_list_committees(client) -> None:
    headers = await _register(client, "pitch-committees@ideaxion.io")
    r = await client.get("/api/v1/pitchsim/committees", headers=headers)
    assert r.status_code == 200, r.text
    keys = {c["key"] for c in r.json()}
    assert keys == {"incubateur", "concours", "investisseur"}
    incub = next(c for c in r.json() if c["key"] == "incubateur")
    assert len(incub["personas"]) == 4


async def test_upload_deck_parses_slides(client) -> None:
    headers = await _register(client, "pitch-deck@ideaxion.io")
    data = _make_pptx(["Probleme", "Solution", "Marche"])

    r = await client.post(
        "/api/v1/pitchsim/decks",
        headers=headers,
        files={"file": ("deck.pptx", data, PPTX_TYPE)},
        data={"title": "Mon pitch"},
    )
    assert r.status_code == 200, r.text
    deck = r.json()
    assert deck["title"] == "Mon pitch"
    assert len(deck["slides"]) == 3
    assert all(s["kind"] == "main" for s in deck["slides"])

    # Relecture du deck par son id.
    r = await client.get(f"/api/v1/pitchsim/decks/{deck['id']}", headers=headers)
    assert r.status_code == 200
    assert len(r.json()["slides"]) == 3


async def test_upload_deck_rejects_bad_type(client) -> None:
    headers = await _register(client, "pitch-badtype@ideaxion.io")
    r = await client.post(
        "/api/v1/pitchsim/decks",
        headers=headers,
        files={"file": ("notes.txt", b"hello", "text/plain")},
        data={"title": "x"},
    )
    assert r.status_code == 422, r.text


async def test_pitch_session_flow_with_committee(client) -> None:
    headers = await _register(client, "pitch-session@ideaxion.io")

    # Démarrer une session devant le comité Incubateur.
    r = await client.post(
        "/api/v1/pitchsim/sessions",
        headers=headers,
        json={"committee_key": "incubateur", "mode": "slides"},
    )
    assert r.status_code == 200, r.text
    ps = r.json()
    assert ps["status"] == "in_progress"
    sid = ps["id"]
    assert any(t["kind"] == "deliberation" for t in ps["turns"])  # briefing

    # Narration faible (courte, sans chiffres) → un juge interrompt sur sa faiblesse.
    r = await client.post(
        f"/api/v1/pitchsim/sessions/{sid}/slide",
        headers=headers,
        json={"narration": "Bonjour, voici mon projet."},
    )
    assert r.status_code == 200, r.text
    turns = r.json()["turns"]
    interruptions = [t for t in turns if t["kind"] == "interruption"]
    assert interruptions, "une faiblesse aurait dû déclencher une interruption"
    assert interruptions[0]["meta"].get("axis")

    # Le porteur répond.
    r = await client.post(
        f"/api/v1/pitchsim/sessions/{sid}/answer",
        headers=headers,
        json={"answer": "Le marché fait 500M€ selon l'étude McKinsey 2025."},
    )
    assert r.status_code == 200

    # Imprévu forcé (entraînement).
    r = await client.post(f"/api/v1/pitchsim/sessions/{sid}/imprevu", headers=headers)
    assert r.status_code == 200
    assert any(t["kind"] == "imprevu" for t in r.json()["turns"])

    # Fin → délibération + statut completed.
    r = await client.post(f"/api/v1/pitchsim/sessions/{sid}/finish", headers=headers)
    assert r.status_code == 200, r.text
    final = r.json()
    assert final["status"] == "completed"
    assert any(t["kind"] == "deliberation" for t in final["turns"])

    # Action après finish interdite (machine à états).
    r = await client.post(
        f"/api/v1/pitchsim/sessions/{sid}/slide",
        headers=headers,
        json={"narration": "encore"},
    )
    assert r.status_code == 422
