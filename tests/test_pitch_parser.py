"""Parsing des decks — pur, avec des fixtures générées en mémoire (PDF + PPTX)."""

from __future__ import annotations

from io import BytesIO

import pytest

from app.pitchsim.parser import parse_deck, parse_pdf, parse_pptx

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _make_pptx(titles: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for t in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # « Title Only »
        slide.shapes.title.text = t
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_blank_pdf(pages: int) -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    buf = BytesIO()
    w.write(buf)
    return buf.getvalue()


def test_parse_pptx_extracts_text_per_slide():
    data = _make_pptx(["Probleme: acces au credit", "Solution: tontine mobile"])
    slides = parse_pptx(data)
    assert len(slides) == 2
    assert "credit" in slides[0]["text"]
    assert slides[0]["title"].startswith("Probleme")


def test_parse_pdf_returns_one_slide_per_page():
    slides = parse_pdf(_make_blank_pdf(3))
    assert len(slides) == 3
    # Pages vierges → texte vide mais titre de repli « Slide N ».
    assert slides[0]["title"] == "Slide 1"


def test_parse_deck_routes_by_content_type():
    assert len(parse_deck(PPTX_TYPE, _make_pptx(["A"]))) == 1
    assert len(parse_deck("application/pdf", _make_blank_pdf(1))) == 1


def test_parse_deck_rejects_unknown_type():
    with pytest.raises(ValueError):
        parse_deck("text/plain", b"hello")
