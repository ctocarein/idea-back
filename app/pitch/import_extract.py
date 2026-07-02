"""Extraction du texte d'un fichier importé (pour générer un deck).

On ne cherche PAS à préserver le design d'origine : on en extrait la MATIÈRE
(le texte), qui repart dans `generate_deck` → un deck propre à notre charte.
Formats : PDF, PPTX, TXT/MD. Imports paresseux (pypdf / python-pptx).
"""

from __future__ import annotations

from io import BytesIO

MAX_CHARS = 12000  # borne : au-delà, on tronque (assez pour un deck)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(run.text for run in para.runs).strip()
                    if txt:
                        parts.append(txt)
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Retourne le texte extrait, borné. Lève ValueError si format non supporté."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _from_pdf(data)
    elif name.endswith(".pptx"):
        text = _from_pptx(data)
    elif name.endswith((".txt", ".md", ".markdown")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Format non supporté (PDF, PPTX, TXT ou MD).")
    text = text.strip()
    if not text:
        raise ValueError("Aucun texte exploitable trouvé dans le fichier.")
    return text[:MAX_CHARS]
