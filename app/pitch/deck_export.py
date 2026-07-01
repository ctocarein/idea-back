"""Export du deck de pitch — PDF et PPTX via Chromium headless (Playwright).

Le deck HTML (`deck_render.render_deck_html(..., export=True)`) est LA source
du design : Playwright le charge dans un vrai navigateur (Chart.js et images
se rendent normalement), puis :
  - PDF  : `page.pdf()`, une page = une slide (960x540, format écran large).
  - PPTX : une capture PNG par slide (`element.screenshot()`), injectée plein
    cadre dans une Presentation `python-pptx` (16:9).

Playwright est optionnel (extra `deck`) ; son absence remonte une ImportError
que l'appelant dégrade en 503 (même contrat que weasyprint/python-pptx).
"""

from __future__ import annotations

from io import BytesIO

SLIDE_WIDTH_PX = 960
SLIDE_HEIGHT_PX = 540


async def _open_deck_page(html: str):
    # Import paresseux : playwright (extra `deck`) n'est chargé que si on
    # exporte réellement. Absence → ImportError, dégradée en 503 par l'appelant.
    from playwright.async_api import async_playwright

    pw = await async_playwright().start()
    browser = await pw.chromium.launch()
    page = await browser.new_page(viewport={"width": SLIDE_WIDTH_PX, "height": SLIDE_HEIGHT_PX})
    await page.set_content(html, wait_until="load")
    # Attend que Chart.js ait fini de dessiner (marqueur posé par deck_render en mode export).
    try:
        await page.wait_for_selector("body[data-deck-ready='1']", timeout=5000)
    except Exception:  # noqa: BLE001 — best-effort, on capture quand même
        pass
    return pw, browser, page


async def render_deck_pdf(html: str) -> bytes:
    """Deck HTML (mode export) → PDF, une page par slide (taille = slide)."""
    pw, browser, page = await _open_deck_page(html)
    try:
        return await page.pdf(
            width=f"{SLIDE_WIDTH_PX}px",
            height=f"{SLIDE_HEIGHT_PX}px",
            print_background=True,
            margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
        )
    finally:
        await browser.close()
        await pw.stop()


async def render_deck_slide_pngs(html: str) -> list[bytes]:
    """Deck HTML (mode export) → une image PNG par slide (pour le PPTX)."""
    pw, browser, page = await _open_deck_page(html)
    try:
        handles = await page.query_selector_all(".slide")
        return [await h.screenshot() for h in handles]
    finally:
        await browser.close()
        await pw.stop()


def build_pptx_from_pngs(pngs: list[bytes]) -> bytes:
    """PNG (un par slide) → PPTX plein cadre, 16:9. Import paresseux python-pptx."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(9144000)  # 10 in
    prs.slide_height = Emu(5143500)  # 5.625 in → 16:9
    blank_layout = prs.slide_layouts[6]

    for png in pngs:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            BytesIO(png), 0, 0, width=prs.slide_width, height=prs.slide_height
        )

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def render_deck_pptx(html: str) -> bytes:
    """Deck HTML (mode export) → PPTX (une slide-image par slide)."""
    pngs = await render_deck_slide_pngs(html)
    return build_pptx_from_pngs(pngs)
