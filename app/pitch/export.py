"""Export du pitch — HTML→PDF (WeasyPrint) et PPTX (python-pptx).

Les fonctions de rendu HTML/PPTX sont pures (testables hors-ligne). Le PDF
dépend de WeasyPrint (extra `pdf`) : son absence remonte une ImportError que
l'appelant dégrade en 503. Idem PPTX si python-pptx est absent.
"""

from __future__ import annotations

from html import escape
from io import BytesIO

from app.pitch.models import Pitch

_INK = "#1C1633"
_CORAL = "#FF7A4D"
_MUTED = "#6B6580"


def _filled_sections(pitch: Pitch) -> list[dict]:
    return [s for s in (pitch.sections or []) if str(s.get("content", "")).strip()]


def render_pitch_html(pitch: Pitch, project_title: str | None = None) -> str:
    title = escape(project_title or pitch.title or "Pitch")
    blocks = []
    for s in _filled_sections(pitch):
        blocks.append(
            f'<section><h2>{escape(str(s.get("title", "")))}</h2>'
            f'<p>{escape(str(s.get("content", ""))).replace(chr(10), "<br>")}</p></section>'
        )
    body = "".join(blocks) or "<p>Ton pitch est encore vide.</p>"
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
  body {{ font-family: "Segoe UI", Roboto, sans-serif; color: {_INK}; font-size: 13px; padding: 8px 0; }}
  h1 {{ color: {_INK}; font-size: 26px; margin: 0 0 4px; }}
  .sub {{ color: {_MUTED}; margin: 0 0 20px; }}
  section {{ margin: 0 0 16px; page-break-inside: avoid; }}
  h2 {{ color: {_CORAL}; font-size: 15px; margin: 0 0 4px; }}
  p {{ line-height: 1.6; margin: 0; }}
  .foot {{ color: {_MUTED}; font-size: 10px; margin-top: 26px; border-top: 1px solid #EEE; padding-top: 8px; }}
</style></head><body>
  <h1>{title}</h1>
  <div class="sub">Pitch · généré avec IDEAXION</div>
  {body}
  <div class="foot">IDEAXION — comprendre, structurer, exprimer ses besoins.</div>
</body></html>"""


def render_pitch_pdf(html: str) -> bytes:
    # WeasyPrint optionnel (extra `pdf`) ; absence → ImportError (dégradée en 503).
    from weasyprint import HTML

    return HTML(string=html).write_pdf()


def render_pitch_pptx(pitch: Pitch, project_title: str | None = None) -> bytes:
    # python-pptx optionnel ; absence → ImportError (dégradée en 503).
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    # Slide de titre.
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = project_title or pitch.title or "Pitch"
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Pitch · généré avec IDEAXION"

    # Une slide par section remplie.
    for s in _filled_sections(pitch):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(s.get("title", ""))
        body = slide.placeholders[1]
        body.text = str(s.get("content", ""))
        for para in body.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(16)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
