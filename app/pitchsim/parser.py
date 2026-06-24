"""Parsing d'un pitch deck → texte par slide (PDF / PPTX).

Pur et sans réseau : prend des octets, rend une liste de slides `{title, text}`. C'est ce
texte qui permet aux juges de commenter une slide précise. Imports paresseux (les libs ne
sont chargées que si on parse réellement). Les vignettes images sont rendues par le front.
"""

from __future__ import annotations

from io import BytesIO

PDF_TYPES = {"application/pdf"}
PPTX_TYPES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
ALLOWED_DECK_TYPES = PDF_TYPES | PPTX_TYPES
MAX_MAIN_SLIDES = 10  # garde-fou : un pitch deck principal reste court


def _title_from(text: str, index: int) -> str:
    # Titre = 1re ligne non vide (tronquée), sinon « Slide N ».
    for line in text.splitlines():
        line = line.strip()
        if line:
            return line[:120]
    return f"Slide {index + 1}"


def parse_pdf(data: bytes) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    slides: list[dict] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        slides.append({"title": _title_from(text, i), "text": text})
    return slides


def parse_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    slides: list[dict] = []
    for i, slide in enumerate(prs.slides):
        parts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        text = "\n".join(parts)
        slides.append({"title": _title_from(text, i), "text": text})
    return slides


def parse_deck(content_type: str, data: bytes) -> list[dict]:
    if content_type in PDF_TYPES:
        return parse_pdf(data)
    if content_type in PPTX_TYPES:
        return parse_pptx(data)
    raise ValueError(f"Type de deck non supporté : {content_type}")
